"""Bulk import service — CSV/XLSX parse, dry-run validation, commit enrollment (T-037/T-079)."""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone

import structlog
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import (
    ConflictError,
    NotFoundError,
    PermissionDeniedError,
    PreconditionFailedError,
    ValidationError,
)
from app.features.academic_sessions.repository import AcademicSessionRepository
from app.features.bulk_imports.models import BulkImport, BulkImportStatus
from app.features.bulk_imports.repository import BulkImportRepository
from app.features.bulk_imports.schemas import BulkImportRead, BulkImportRowResult
from app.features.files.pipeline import run_upload_pipeline
from app.features.files.profiles import get_profile
from app.features.grades.repository import GradeRepository
from app.features.independent_users.repository import IndependentUserRepository
from app.features.invites.repository import UserInviteRepository
from app.features.sections.repository import SectionRepository
from app.features.student_enrollments.schemas import StudentEnrollmentCreate
from app.features.student_enrollments.service import StudentEnrollmentService
from app.features.users.models import User, UserRole
from app.features.users.repository import UserRepository
from app.infrastructure.audit.log import audit
from app.infrastructure.notifications.account import notify_account_event

logger = structlog.get_logger(__name__)

_PROFILE_NAME = "bulk_import"
MAX_ROWS = 5000
REQUIRED_COLUMNS = frozenset({"name", "email", "grade"})
OPTIONAL_COLUMNS = frozenset({"section", "language"})
ALLOWED_LANGUAGES = frozenset({"en", "ur", "sd", "ps"})
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class BulkImportService:
    """Parse upload files, validate rows, and commit student enrollments."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session
        self._repo = BulkImportRepository(session)
        self._users = UserRepository(session)
        self._invites = UserInviteRepository(session)
        self._independent_users = IndependentUserRepository(session)
        self._grades = GradeRepository(session)
        self._sections = SectionRepository(session)
        self._sessions = AcademicSessionRepository(session)

    async def _require_coordinator(self, actor_authentik_id: str) -> User:
        actor = await self._users.get_by_authentik_id(actor_authentik_id)
        if actor is None or actor.role != UserRole.COORDINATOR:
            raise PermissionDeniedError("Only coordinators can run bulk imports")
        if not actor.school_id:
            raise ValidationError("Coordinator must belong to a school")
        grade_scope = _parse_grade_scope(actor.scoped_ids)
        if not grade_scope:
            raise ValidationError("Coordinator has no grade scope configured")
        return actor

    async def _active_session_label(self, school_id: str) -> str | None:
        label = await self._sessions.get_school_active_label(school_id)
        if not label:
            active = await self._sessions.get_active(school_id)
            label = active.label if active else None
        return label

    async def dry_run(
        self,
        data: bytes,
        filename: str,
        actor_authentik_id: str,
    ) -> BulkImportRead:
        """Validate rows and persist a dry-run job (no user creation)."""
        actor = await self._require_coordinator(actor_authentik_id)
        grade_scope = _parse_grade_scope(actor.scoped_ids)

        rows_raw = _parse_file(data, filename)
        if not rows_raw:
            raise ValidationError("File contains no data rows")
        if len(rows_raw) > MAX_ROWS:
            raise ValidationError(f"File exceeds maximum of {MAX_ROWS} rows")

        session_label = await self._active_session_label(actor.school_id)
        row_results = await self._validate_rows(
            rows_raw, grade_scope, actor.school_id, session_label
        )
        success_rows = sum(1 for r in row_results if r.status == "valid")
        failed_rows = len(row_results) - success_rows

        profile = get_profile(_PROFILE_NAME)
        is_csv = _is_csv(filename, data)
        upload_result = await run_upload_pipeline(
            data=data,
            filename=filename,
            profile=profile,
            session=self._session,
            school_id=actor.school_id,
            uploaded_by=actor.id,
            skip_magic_check=is_csv,
        )

        now = datetime.now(timezone.utc)
        job = BulkImport(
            school_id=actor.school_id,
            imported_by_user_id=actor.id,
            upload_id=upload_result.upload_id,
            total_rows=len(row_results),
            success_rows=success_rows,
            failed_rows=failed_rows,
            error_report_jsonb={"rows": [r.model_dump() for r in row_results]},
            status=BulkImportStatus.DRY_RUN_COMPLETE,
            completed_at=now,
        )
        job = await self._repo.save(job)

        await audit(
            session=self._session,
            action="bulk_import.dry_run_complete",
            actor_id=actor.id,
            target_type="bulk_import",
            target_id=job.id,
            metadata={
                "total_rows": job.total_rows,
                "success_rows": job.success_rows,
                "failed_rows": job.failed_rows,
            },
        )

        logger.info(
            "bulk_import_dry_run_complete",
            import_id=job.id,
            school_id=actor.school_id,
            total=job.total_rows,
            valid=job.success_rows,
        )
        return _to_read(job)

    async def commit(
        self,
        import_id: str,
        actor_authentik_id: str,
        claims: dict[str, object],
    ) -> BulkImportRead:
        """Enroll all dry-run-valid rows; partial success allowed."""
        actor = await self._require_coordinator(actor_authentik_id)
        job = await self._repo.get_by_id_for_school(import_id, actor.school_id)
        if job is None:
            raise NotFoundError("Bulk import not found")
        if job.status != BulkImportStatus.DRY_RUN_COMPLETE:
            raise PreconditionFailedError("Bulk import is not ready to commit")

        report = job.error_report_jsonb or {}
        row_results = [BulkImportRowResult.model_validate(r) for r in report.get("rows", [])]
        enrollment_svc = StudentEnrollmentService(self._session)
        actor_id = str(claims.get("sub", ""))

        for row in row_results:
            if row.status != "valid" or not row.data:
                continue
            grade_id = row.data.get("grade_id", "").strip()
            if not grade_id:
                row.status = "failed"
                row.errors = [*row.errors, "grade_not_found"]
                continue

            payload = StudentEnrollmentCreate(
                display_name=row.data["name"],
                email=row.data["email"],
                section_id=row.data.get("section_id") or None,
            )
            try:
                await enrollment_svc.enroll_student(grade_id, payload, claims, actor_id)
                row.status = "enrolled"
            except ConflictError:
                row.status = "failed"
                row.errors = [*row.errors, "email_already_exists"]
            except PermissionDeniedError:
                row.status = "failed"
                row.errors = [*row.errors, "grade_out_of_scope"]
            except (ValidationError, NotFoundError):
                row.status = "failed"
                row.errors = [*row.errors, "enrollment_failed"]

        enrolled = sum(1 for r in row_results if r.status == "enrolled")
        job.error_report_jsonb = {"rows": [r.model_dump() for r in row_results]}
        job.success_rows = enrolled
        job.failed_rows = job.total_rows - enrolled
        job.status = (
            BulkImportStatus.COMMITTED
            if job.failed_rows == 0
            else BulkImportStatus.COMMITTED_WITH_ERRORS
        )
        job.completed_at = datetime.now(timezone.utc)
        job = await self._repo.update(job)

        await audit(
            session=self._session,
            action="bulk_import.committed",
            actor_id=actor.id,
            target_type="bulk_import",
            target_id=job.id,
            school_id=actor.school_id,
            metadata={
                "total_rows": job.total_rows,
                "enrolled_rows": enrolled,
                "failed_rows": job.failed_rows,
            },
        )
        await notify_account_event(
            session=self._session,
            template_key="account.bulk_import_done",
            recipient_user_id=actor.id,
            school_id=actor.school_id,
            params={
                "total_rows": str(job.total_rows),
                "success_rows": str(job.success_rows),
                "failed_rows": str(job.failed_rows),
            },
            metadata={"bulk_import_id": job.id},
        )
        logger.info(
            "bulk_import_committed",
            import_id=job.id,
            enrolled=enrolled,
            failed=job.failed_rows,
        )
        return _to_read(job)

    async def get_job(self, import_id: str, actor_authentik_id: str) -> BulkImportRead:
        """Return a bulk import job scoped to the coordinator's school."""
        actor = await self._require_coordinator(actor_authentik_id)
        job = await self._repo.get_by_id_for_school(import_id, actor.school_id)
        if job is None:
            raise NotFoundError("Bulk import not found")
        return _to_read(job)

    async def _validate_rows(
        self,
        rows_raw: list[dict[str, str]],
        grade_scope: set[str],
        school_id: str,
        session_label: str | None,
    ) -> list[BulkImportRowResult]:
        seen_emails: set[str] = set()
        results: list[BulkImportRowResult] = []

        for index, raw in enumerate(rows_raw, start=2):
            errors: list[str] = []
            name = raw.get("name", "").strip()
            email = raw.get("email", "").strip().lower()
            grade = raw.get("grade", "").strip()
            section = raw.get("section", "").strip()
            language = raw.get("language", "").strip().lower() or "en"
            grade_id = ""
            section_id = ""

            if not name:
                errors.append("missing_name")
            if not email:
                errors.append("missing_email")
            elif not _EMAIL_RE.match(email):
                errors.append("invalid_email")
            elif email in seen_emails:
                errors.append("duplicate_email_in_file")
            else:
                seen_emails.add(email)
                existing = await self._users.get_by_email(email)
                if existing is not None:
                    if existing.school_id != school_id:
                        errors.append("email_in_other_school")
                    else:
                        errors.append("email_already_exists")
                elif await self._invites.get_pending_by_email(email) is not None:
                    errors.append("email_already_exists")
                elif await self._independent_users.get_by_email(email) is not None:
                    errors.append("email_already_exists")

            if not grade:
                errors.append("missing_grade")
            elif grade not in grade_scope:
                errors.append("grade_out_of_scope")
            elif not session_label:
                errors.append("grade_not_found")
            else:
                grade_row = await self._grades.get_by_name_session(
                    school_id, grade, session_label
                )
                if grade_row is None:
                    errors.append("grade_not_found")
                else:
                    grade_id = grade_row.id
                    if section:
                        sec = await self._sections.get_by_name(grade_row.id, section)
                        if sec is None:
                            errors.append("section_not_found")
                        else:
                            section_id = sec.id
                    else:
                        sections = await self._sections.list_by_grade(grade_row.id)
                        default = next((s for s in sections if s.is_default_internal), None)
                        if default is None:
                            errors.append("section_not_found")
                        else:
                            section_id = default.id

            if language and language not in ALLOWED_LANGUAGES:
                errors.append("invalid_language")

            data = {
                "name": name,
                "email": email,
                "grade": grade,
                "section": section,
                "language": language,
                "grade_id": grade_id,
                "section_id": section_id,
            }
            results.append(
                BulkImportRowResult(
                    row_number=index,
                    status="valid" if not errors else "invalid",
                    errors=errors,
                    data=data if not errors else data,
                )
            )

        return results


def _parse_grade_scope(scoped_ids: str | None) -> set[str]:
    if not scoped_ids:
        return set()
    return {part.strip() for part in scoped_ids.split(",") if part.strip()}


def _is_csv(filename: str, data: bytes) -> bool:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return True
    return not data.startswith(b"PK\x03\x04")


def _normalize_header(header: object) -> str:
    return str(header or "").strip().lower()


def _parse_file(data: bytes, filename: str) -> list[dict[str, str]]:
    if _is_csv(filename, data):
        return _parse_csv(data)
    return _parse_xlsx(data)


def _parse_csv(data: bytes) -> list[dict[str, str]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise ValidationError("CSV file is missing a header row")

    normalized_fields = {_normalize_header(h): h for h in reader.fieldnames}
    missing = REQUIRED_COLUMNS - set(normalized_fields)
    if missing:
        raise ValidationError(f"Missing required columns: {', '.join(sorted(missing))}")

    rows: list[dict[str, str]] = []
    for raw_row in reader:
        if not any(str(v or "").strip() for v in raw_row.values()):
            continue
        row: dict[str, str] = {}
        for col in REQUIRED_COLUMNS | OPTIONAL_COLUMNS:
            source_key = normalized_fields.get(col)
            row[col] = str(raw_row.get(source_key, "") if source_key else "").strip()
        rows.append(row)
    return rows


def _parse_xlsx(data: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValidationError("Spreadsheet is empty") from exc

    headers = [_normalize_header(cell) for cell in header_row]
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise ValidationError(f"Missing required columns: {', '.join(sorted(missing))}")

    col_index = {name: headers.index(name) for name in REQUIRED_COLUMNS | OPTIONAL_COLUMNS if name in headers}
    rows: list[dict[str, str]] = []
    for raw in rows_iter:
        if raw is None or not any(str(cell or "").strip() for cell in raw):
            continue
        row: dict[str, str] = {}
        for col in REQUIRED_COLUMNS | OPTIONAL_COLUMNS:
            idx = col_index.get(col)
            value = raw[idx] if idx is not None and idx < len(raw) else ""
            row[col] = str(value or "").strip()
        rows.append(row)
    return rows


def _to_read(job: BulkImport) -> BulkImportRead:
    report = job.error_report_jsonb or {}
    raw_rows = report.get("rows", [])
    rows = [BulkImportRowResult.model_validate(r) for r in raw_rows]
    return BulkImportRead(
        id=job.id,
        school_id=job.school_id,
        imported_by_user_id=job.imported_by_user_id,
        upload_id=job.upload_id,
        total_rows=job.total_rows,
        success_rows=job.success_rows,
        failed_rows=job.failed_rows,
        status=job.status.value,
        rows=rows,
        created_at=job.created_at,
        completed_at=job.completed_at,  # type: ignore[arg-type]
    )
